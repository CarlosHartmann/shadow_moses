import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


SCRIPT_DIR = Path(__file__).resolve().parents[1] / "0-assets" / "raw_data" / "main" / "3b-age-grouping"
sys.path.insert(0, str(SCRIPT_DIR))

import group_by_age
import group_by_generation


def read_ids(path: Path) -> list[str]:
    return [json.loads(line)["id"] for line in path.read_text(encoding="utf-8").splitlines()]


class AgeSplitTests(unittest.TestCase):
    def test_generation_majority(self):
        classify = group_by_generation.classify
        self.assertEqual(classify([1990, 1991]), "millennial")
        self.assertEqual(classify([1975, 1985]), "X")  # 6 of 11 years
        self.assertEqual(classify([1970, 1981]), "X")  # 11 of 12 years
        self.assertEqual(classify([1996, 1997]), "millennial")  # tie goes to older
        self.assertEqual(classify([2020, 2021]), None)
        self.assertEqual(classify(None), None)

    def test_young_old_cutoff(self):
        classify = group_by_age.classify
        self.assertEqual(classify([1980, 1980]), "old")
        self.assertEqual(classify([1981, 1981]), "young")
        self.assertEqual(classify([1979, 1981]), "old")
        self.assertEqual(classify([1980, 1981]), "old")  # tie goes to older

    def test_unassigned_comments_are_dropped(self):
        comments = [
            {"id": "a", "created_utc": "1579000000", "deduced_birthyear_range": [1990, 1991]},
            {"id": "b", "created_utc": "1579000000", "deduced_birthyear_range": [1960, 1961]},
            {"id": "c", "created_utc": "1579000000", "deduced_birthyear_range": None},
        ]
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "age.jsonl"
            source.write_text("".join(json.dumps(item) + "\n" for item in comments), encoding="utf-8")
            group_by_age.split_by_bucket(
                source, root / "out", group_by_age.FILE_NAMES, group_by_age.classify
            )
            self.assertEqual(read_ids(root / "out" / "age_young.jsonl"), ["a"])
            self.assertEqual(read_ids(root / "out" / "age_old.jsonl"), ["b"])
            self.assertEqual(sorted(path.name for path in (root / "out").glob("age_*.jsonl")), ["age_old.jsonl", "age_young.jsonl"])

    def test_quarterly_statistics_count_written_comments(self):
        january_2020 = "1579000000"
        comments = [
            {"id": "a", "created_utc": january_2020, "deduced_birthyear_range": [1990, 1991]},
            {"id": "b", "created_utc": january_2020, "deduced_birthyear_range": [1960, 1961]},
            {"id": "c", "created_utc": january_2020, "deduced_birthyear_range": None},
        ]
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "age.jsonl"
            source.write_text("".join(json.dumps(item) + "\n" for item in comments), encoding="utf-8")
            group_by_age.split_by_bucket(
                source, root / "out", group_by_age.FILE_NAMES, group_by_age.classify, "stats.xlsx"
            )
            workbook = load_workbook(root / "out" / "stats.xlsx", read_only=True)
            rows = list(workbook["Quarterly counts"].iter_rows(values_only=True))
            workbook.close()
        self.assertEqual(rows[0], ("quarter", "young", "old"))
        self.assertEqual(rows[1][0], "2010-Q1")
        self.assertEqual(next(row for row in rows if row[0] == "2020-Q1"), ("2020-Q1", 1, 1))


if __name__ == "__main__":
    unittest.main()
