import json
import sys
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


SCRIPT_DIR = (
    Path(__file__).resolve().parents[1]
    / "0-assets"
    / "raw_data"
    / "main"
    / "4-cross-subreddit-data"
)
sys.path.insert(0, str(SCRIPT_DIR))

from cross_subreddit import (
    DEFAULT_OUTPUT_DIR,
    GROUPS,
    group_path,
    process_cross_subreddit_data,
)


def make_comment(
    identifier: str, author: str, subreddit: str, year: int = 2020
) -> dict:
    timestamp = datetime(year, 1, 15, tzinfo=timezone.utc).timestamp()
    return {
        "id": identifier,
        "author": author,
        "subreddit": subreddit,
        "created_utc": str(timestamp),
    }


def write_jsonl(path: Path, comments: list[dict]) -> None:
    with path.open("w", encoding="utf-8") as output_file:
        for item in comments:
            output_file.write(json.dumps(item) + "\n")


class CrossSubredditTests(unittest.TestCase):
    def test_default_output_dir_is_dedicated_subdirectory(self) -> None:
        self.assertEqual(DEFAULT_OUTPUT_DIR, SCRIPT_DIR / "cross_subreddit_data")

    def test_partitions_baseline_and_adds_unseen_comments_to_author_groups(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            grouped_dir = root / "grouped"
            age_dir = root / "age"
            output_dir = root / "output"
            grouped_dir.mkdir()
            age_dir.mkdir()
            original = make_comment("original", "Alice", "study_subreddit")
            pre_2012_original = make_comment(
                "pre-2012-original", "PreOnly", "study_subreddit", year=2011
            )
            for group in GROUPS:
                write_jsonl(group_path(group, grouped_dir, age_dir), [])
            write_jsonl(age_dir / "age_millennial.jsonl", [original])
            write_jsonl(age_dir / "age_young.jsonl", [original])
            write_jsonl(age_dir / "age_old.jsonl", [pre_2012_original])
            write_jsonl(grouped_dir / "gender_male.jsonl", [original])

            baseline_path = root / "baseline.jsonl"
            write_jsonl(
                baseline_path,
                [
                    make_comment("baseline-extra", "ALICE", "outside_subreddits"),
                    make_comment("pre-only-baseline", "PreOnly", "outside_subreddits"),
                    make_comment("unrelated", "Carol", "outside_subreddits"),
                ],
            )
            filtered_path = root / "filtered.jsonl"
            write_jsonl(
                filtered_path,
                [
                    original,
                    pre_2012_original,
                    make_comment("baseline-extra", "Alice", "outside_subreddits"),
                    make_comment("alice-pre-2012", "Alice", "historical_subreddit", year=2011),
                    make_comment("pre-only-new", "PreOnly", "elsewhere"),
                    make_comment("unseen-target", "Alice", "study_subreddit"),
                    make_comment("unseen-other", "Alice", "elsewhere"),
                    make_comment("unrelated", "Carol", "outside_subreddits"),
                ],
            )

            process_cross_subreddit_data(
                baseline_path, grouped_dir, filtered_path, output_dir, age_dir
            )

            expected_ids = [
                "baseline-extra",
                "alice-pre-2012",
                "unseen-target",
                "unseen-other",
            ]
            for group in ("age_millennial", "age_young", "gender_male"):
                with (output_dir / f"{group}.jsonl").open(encoding="utf-8") as group_file:
                    identifiers = [json.loads(line)["id"] for line in group_file]
                self.assertEqual(identifiers, expected_ids)
            with (output_dir / "age_old.jsonl").open(encoding="utf-8") as group_file:
                self.assertEqual([json.loads(line)["id"] for line in group_file], [])
            self.assertFalse((output_dir / "age.jsonl").exists())

            with (output_dir / "baseline.jsonl").open(encoding="utf-8") as baseline_file:
                remaining_ids = [json.loads(line)["id"] for line in baseline_file]
            self.assertEqual(remaining_ids, ["pre-only-baseline", "unrelated"])

            workbook = load_workbook(
                output_dir / "quarterly_statistics.xlsx", read_only=True
            )
            rows = list(workbook["Quarterly counts"].iter_rows(values_only=True))
            self.assertNotIn("Warnings", workbook.sheetnames)
            workbook.close()
            self.assertEqual(rows[0][0], "quarter")
            column = {name: index for index, name in enumerate(rows[0])}
            self.assertNotIn("age", column)
            quarter_counts = next(row for row in rows[1:] if row[0] == "2020-Q1")
            self.assertEqual(quarter_counts[column["baseline"]], 2)
            self.assertEqual(quarter_counts[column["age_millennial"]], 4)
            self.assertEqual(quarter_counts[column["age_young"]], 4)
            self.assertEqual(quarter_counts[column["age_old"]], 0)
            self.assertEqual(quarter_counts[column["gender_male"]], 4)
            historical_counts = next(row for row in rows[1:] if row[0] == "2011-Q1")
            self.assertEqual(historical_counts[column["baseline"]], 0)
            self.assertEqual(historical_counts[column["age_millennial"]], 1)
            self.assertEqual(historical_counts[column["age_old"]], 1)
            self.assertEqual(historical_counts[column["gender_male"]], 1)


if __name__ == "__main__":
    unittest.main()