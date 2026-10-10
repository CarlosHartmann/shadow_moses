"""Add `deduced_birthyear_range` to the age comments, based on the annotated "age" sheet of grouped_flairs.xlsx."""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent
GROUPING_DIR = HERE.parent / "3a-grouping"
GROUPED_FLAIRS_PATH = GROUPING_DIR / "grouped_flairs.xlsx"
SHEET_NAME = "age"
DEFAULT_INPUT = GROUPING_DIR / "grouped_subreddit_data" / "age.jsonl"
DEFAULT_OUTPUT = HERE / "age.jsonl"
OUTPUT_FIELD = "deduced_birthyear_range"

# xennial is not standardized; the other ranges follow Pew-style definitions.
GENERATION_RANGES = {
    "silent": (1928, 1945),
    "boomer": (1946, 1964),
    "x": (1965, 1980),
    "xennial": (1975, 1985),
    "millennial": (1981, 1996),
    "z": (1997, 2012),
}
RANGE_PATTERN = re.compile(r"^\s*(\d{1,4})\s*-\s*(\d{1,4})\s*$")


def parse_range(text) -> tuple[int, int] | None:
    match = RANGE_PATTERN.match(str(text)) if text not in (None, "") else None
    if not match:
        return None
    low, high = sorted((int(match[1]), int(match[2])))
    return low, high


def parse_int(value) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and float(value).is_integer():
        return int(value)
    if isinstance(value, str) and value.strip().isdigit():
        return int(value.strip())
    return None


def birth_range_for(annotation: dict, year: int) -> tuple[int, int] | None:
    """Priority: birthyear, birthyear range, age, age range, generation."""
    birthyear = parse_int(annotation.get("given_birthyear"))
    if birthyear is not None:
        return birthyear, birthyear
    birthyear_range = parse_range(annotation.get("given_birthyear_range"))
    if birthyear_range:
        return birthyear_range
    age = parse_int(annotation.get("given_age"))
    if age is not None:
        # The birthday may or may not have passed yet in the comment year.
        return year - age - 1, year - age
    age_range = parse_range(annotation.get("given_agerange"))
    if age_range:
        return year - age_range[1] - 1, year - age_range[0]
    generation = str(annotation.get("given_generation") or "").strip().casefold()
    return GENERATION_RANGES.get(generation)


def load_annotations(path: Path) -> tuple[dict[tuple[str, str], dict], Counter]:
    """Map (subreddit, flair) to annotation columns; a flair can occupy several rows, so rows are merged."""
    workbook = load_workbook(path, read_only=True)
    rows = workbook[SHEET_NAME].iter_rows(values_only=True)
    header = [str(name) for name in next(rows)]
    annotations = {}
    problems = Counter()
    for row in rows:
        record = dict(zip(header, row))
        key = (str(record["subreddit"] or "").casefold(), str(record["flair"] or "").strip())
        if not key[1]:
            continue
        generation = str(record.get("given_generation") or "").strip().casefold()
        if generation and generation not in GENERATION_RANGES:
            problems[f"unknown generation label {generation!r}"] += 1
            record["given_generation"] = None
        for column in ("given_age", "given_birthyear"):
            if record.get(column) not in (None, "") and parse_int(record[column]) is None:
                problems[f"non-numeric {column}"] += 1
        for column in ("given_agerange", "given_birthyear_range"):
            if record.get(column) not in (None, "") and parse_range(record[column]) is None:
                problems[f"unparseable {column}"] += 1
        merged = annotations.setdefault(key, {})
        for column, value in record.items():
            if value in (None, ""):
                continue
            if merged.get(column) not in (None, "") and merged[column] != value:
                problems[f"conflicting duplicate rows in {column}"] += 1
                continue
            merged[column] = value
    workbook.close()
    return annotations, problems


def comment_year(comment: dict) -> int:
    return datetime.fromtimestamp(float(comment["created_utc"]), tz=timezone.utc).year


def annotate_file(input_path: Path, output_path: Path, annotations: dict) -> Counter:
    stats = Counter()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with input_path.open("r", encoding="utf-8") as source, output_path.open(
        "w", encoding="utf-8"
    ) as target:
        for line in source:
            if not line.strip():
                continue
            comment = json.loads(line)
            key = (
                str(comment.get("subreddit", "")).casefold(),
                str(comment.get("author_flair_text") or "").strip(),
            )
            annotation = annotations.get(key)
            result = birth_range_for(annotation, comment_year(comment)) if annotation else None
            comment[OUTPUT_FIELD] = list(result) if result else None
            stats["comments"] += 1
            stats["annotated" if result else "no birth year"] += 1
            target.write(json.dumps(comment, ensure_ascii=False) + "\n")
    return stats


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--flairs", type=Path, default=GROUPED_FLAIRS_PATH)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    annotations, problems = load_annotations(args.flairs)
    for message, count in problems.items():
        print(f"warning: {message} in {count} rows (ignored)")
    stats = annotate_file(args.input, args.output, annotations)
    print(
        f"{args.input.name}: {stats['comments']:,} comments, "
        f"{stats['annotated']:,} with {OUTPUT_FIELD}, {stats['no birth year']:,} without"
    )


if __name__ == "__main__":
    main()
