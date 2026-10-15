"""Collect additional comments from authors identified in the grouped data."""

from __future__ import annotations

import argparse
import json
import time
from collections import defaultdict
from contextlib import ExitStack
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterator, Mapping

from openpyxl import Workbook, load_workbook


HERE = Path(__file__).resolve().parent
GROUPED_DIR = HERE.parent / "3a-grouping" / "grouped_subreddit_data"
AGE_DIR = HERE.parent / "3b-age-grouping"
# Files written by the 3b group_by scripts; both splits cover the same comments.
AGE_GROUPS = (
    "age_silent",
    "age_boomer",
    "age_X",
    "age_millennial",
    "age_Z",
    "age_young",
    "age_old",
)
DEFAULT_BASELINE = GROUPED_DIR / "baseline.jsonl"
GROUPED_FLAIRS_PATH = HERE.parent / "3a-grouping" / "grouped_flairs.xlsx"
DEFAULT_FILTERED_INPUT = HERE.parent / "2-filtering" / "filtered_comments.jsonl"
DEFAULT_OUTPUT_DIR = HERE / "cross_subreddit_data"
STATISTICS_START_YEAR = 2010
AUTHOR_SELECTION_START = datetime(2012, 1, 1, tzinfo=timezone.utc).timestamp()
GROUPS = (
    *AGE_GROUPS,
    "gender_male",
    "gender_female_or_diverse",
    "political_conservative",
    "political_progressive",
    "gender_political_conservative",
    "gender_political_progressive",
)
STATISTICS_GROUPS = ("baseline", *GROUPS)


def iter_comments(path: Path) -> Iterator[tuple[int, str, dict]]:
    with path.open("r", encoding="utf-8") as input_file:
        for line_number, line in enumerate(input_file, start=1):
            if not line.strip():
                continue
            try:
                comment = json.loads(line)
            except json.JSONDecodeError as error:
                raise ValueError(f"Invalid JSON on line {line_number} of {path}: {error}") from error
            if not isinstance(comment, dict):
                raise ValueError(f"Expected a JSON object on line {line_number} of {path}.")
            yield line_number, line if line.endswith("\n") else line + "\n", comment


def author_key(comment: dict) -> str:
    author = str(comment.get("author") or "").strip()
    if not author or author.casefold() in {"[deleted]", "[removed]"}:
        return ""
    return author.casefold()


def comment_id(comment: dict, line_number: int, path: Path) -> str:
    value = str(comment.get("id") or "").strip()
    if not value:
        raise ValueError(f"Missing comment id on line {line_number} of {path}.")
    return value


def quarter_for(comment: dict, line_number: int, path: Path) -> str:
    try:
        timestamp = float(comment["created_utc"])
        created_at = datetime.fromtimestamp(timestamp, tz=timezone.utc)
    except (KeyError, TypeError, ValueError, OverflowError, OSError) as error:
        raise ValueError(
            f"Missing or invalid created_utc on line {line_number} of {path}."
        ) from error
    quarter = (created_at.month - 1) // 3 + 1
    return f"{created_at.year}-Q{quarter}"


def quarter_sequence(first: str, last: str) -> list[str]:
    year, quarter = map(int, first.replace("-Q", " ").split())
    last_year, last_quarter = map(int, last.replace("-Q", " ").split())
    quarters = []
    while (year, quarter) <= (last_year, last_quarter):
        quarters.append(f"{year}-Q{quarter}")
        year, quarter = (year + 1, 1) if quarter == 4 else (year, quarter + 1)
    return quarters


def group_path(group: str, grouped_dir: Path, age_dir: Path) -> Path:
    return (age_dir if group in AGE_GROUPS else grouped_dir) / f"{group}.jsonl"


def load_grouped_users_and_ids(
    grouped_dir: Path,
    age_dir: Path,
) -> tuple[
    dict[str, set[str]],
    set[str],
    defaultdict[str, defaultdict[str, int]],
    defaultdict[str, defaultdict[str, set[str]]],
    str | None,
]:
    groups_by_author: dict[str, set[str]] = defaultdict(set)
    original_ids: set[str] = set()
    counts: defaultdict[str, defaultdict[str, int]] = defaultdict(
        lambda: defaultdict(int)
    )
    distinct_users: defaultdict[str, defaultdict[str, set[str]]] = defaultdict(
        lambda: defaultdict(set)
    )
    last_quarter = None

    for group in GROUPS:
        path = group_path(group, grouped_dir, age_dir)
        for line_number, _, comment in iter_comments(path):
            quarter = quarter_for(comment, line_number, path)
            author = author_key(comment)
            if author and float(comment["created_utc"]) >= AUTHOR_SELECTION_START:
                groups_by_author[author].add(group)
            original_ids.add(comment_id(comment, line_number, path))
            if int(quarter[:4]) >= STATISTICS_START_YEAR:
                counts[quarter][group] += 1
                if author:
                    distinct_users[quarter][group].add(author)
                if last_quarter is None or quarter > last_quarter:
                    last_quarter = quarter

    return groups_by_author, original_ids, counts, distinct_users, last_quarter


def load_deleted_authors(path: Path) -> set[str]:
    """Users deleted from the gender or age groups; they belong in no group and not in the baseline."""
    if not path.exists():
        return set()
    workbook = load_workbook(path, read_only=True)
    deleted: set[str] = set()
    if "user_group_overrides" in workbook.sheetnames:
        for user, family, keep in workbook["user_group_overrides"].iter_rows(
            min_row=2, max_col=3, values_only=True
        ):
            if user and family in ("gender", "age") and keep == "none":
                deleted.add(str(user).casefold())
    workbook.close()
    return deleted


def overlap_group_name(groups: set[str]) -> str | None:
    if len(groups) < 2:
        return None
    return " + ".join(group for group in GROUPS if group in groups)


def load_overlap_statistics(
    grouped_dir: Path,
    age_dir: Path,
    groups_by_author: Mapping[str, set[str]],
) -> tuple[
    defaultdict[str, defaultdict[str, int]],
    defaultdict[str, defaultdict[str, set[str]]],
]:
    comment_counts: defaultdict[str, defaultdict[str, int]] = defaultdict(
        lambda: defaultdict(int)
    )
    distinct_users: defaultdict[str, defaultdict[str, set[str]]] = defaultdict(
        lambda: defaultdict(set)
    )
    counted_comment_ids: set[str] = set()

    for group in GROUPS:
        path = group_path(group, grouped_dir, age_dir)
        for line_number, _, comment in iter_comments(path):
            author = author_key(comment)
            overlap = overlap_group_name(groups_by_author.get(author, set()))
            if overlap is None:
                continue
            quarter = quarter_for(comment, line_number, path)
            if int(quarter[:4]) < STATISTICS_START_YEAR:
                continue
            distinct_users[quarter][overlap].add(author)
            identifier = comment_id(comment, line_number, path)
            if identifier not in counted_comment_ids:
                counted_comment_ids.add(identifier)
                comment_counts[quarter][overlap] += 1

    return comment_counts, distinct_users


def write_statistics(
    output_path: Path,
    quarters: list[str],
    counts: Mapping[str, Mapping[str, int]],
    distinct_users: Mapping[str, Mapping[str, set[str]]],
    overlap_comment_counts: Mapping[str, Mapping[str, int]],
    overlap_distinct_users: Mapping[str, Mapping[str, set[str]]],
    run_info: list[tuple[str, str]],
) -> None:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Quarterly counts")
    sheet.append(["quarter", *STATISTICS_GROUPS])
    for quarter in quarters:
        quarter_counts = counts.get(quarter, {})
        sheet.append(
            [quarter, *(quarter_counts.get(group, 0) for group in STATISTICS_GROUPS)]
        )
    user_sheet = workbook.create_sheet("Distinct users")
    user_sheet.append(["quarter", *STATISTICS_GROUPS])
    for quarter in quarters:
        quarter_users = distinct_users.get(quarter, {})
        user_sheet.append(
            [quarter, *(len(quarter_users.get(group, set())) for group in STATISTICS_GROUPS)]
        )
    overlap_groups = sorted(
        {
            overlap
            for quarter_counts in overlap_comment_counts.values()
            for overlap in quarter_counts
        }
        | {
            overlap
            for quarter_users in overlap_distinct_users.values()
            for overlap in quarter_users
        }
    )
    overlap_sheet = workbook.create_sheet("Overlap comments")
    overlap_sheet.append(["quarter", *overlap_groups])
    for quarter in quarters:
        quarter_counts = overlap_comment_counts.get(quarter, {})
        overlap_sheet.append(
            [quarter, *(quarter_counts.get(overlap, 0) for overlap in overlap_groups)]
        )
    overlap_user_sheet = workbook.create_sheet("Overlap users")
    overlap_user_sheet.append(["quarter", *overlap_groups])
    for quarter in quarters:
        quarter_users = overlap_distinct_users.get(quarter, {})
        overlap_user_sheet.append(
            [quarter, *(len(quarter_users.get(overlap, set())) for overlap in overlap_groups)]
        )
    info_sheet = workbook.create_sheet("Run info")
    for row in run_info:
        info_sheet.append(list(row))
    workbook.save(output_path)


def process_cross_subreddit_data(
    baseline_path: Path,
    grouped_dir: Path,
    filtered_path: Path,
    output_dir: Path,
    age_dir: Path = AGE_DIR,
    excluded_authors: set[str] | None = None,
) -> None:
    run_started = time.monotonic()
    run_started_at = datetime.now().isoformat(timespec="seconds")
    output_dir.mkdir(parents=True, exist_ok=True)
    (
        groups_by_author,
        original_ids,
        counts,
        distinct_users,
        last_quarter,
    ) = load_grouped_users_and_ids(grouped_dir, age_dir)
    overlap_comment_counts, overlap_distinct_users = load_overlap_statistics(
        grouped_dir, age_dir, groups_by_author
    )
    baseline_spill = output_dir / ".baseline.spill.jsonl"
    group_spills = {group: output_dir / f".{group}.spill.jsonl" for group in GROUPS}
    baseline_comments = 0
    baseline_removed = 0
    filtered_comments = 0
    cross_comment_ids: set[str] = set()
    cross_group_counts = defaultdict(int)

    try:
        with ExitStack() as stack:
            baseline_output = stack.enter_context(
                baseline_spill.open("w", encoding="utf-8")
            )
            group_outputs = {
                group: stack.enter_context(path.open("w", encoding="utf-8"))
                for group, path in group_spills.items()
            }

            for line_number, line, comment in iter_comments(baseline_path):
                baseline_comments += 1
                author = author_key(comment)
                if author in groups_by_author or author in (excluded_authors or ()):
                    baseline_removed += 1
                    continue
                baseline_output.write(line)
                quarter = quarter_for(comment, line_number, baseline_path)
                if int(quarter[:4]) >= STATISTICS_START_YEAR:
                    counts[quarter]["baseline"] += 1
                    if author:
                        distinct_users[quarter]["baseline"].add(author)
                    if last_quarter is None or quarter > last_quarter:
                        last_quarter = quarter

            for line_number, line, comment in iter_comments(filtered_path):
                filtered_comments += 1
                author = author_key(comment)
                groups = groups_by_author.get(author)
                if not groups:
                    continue
                identifier = comment_id(comment, line_number, filtered_path)
                if identifier in original_ids or identifier in cross_comment_ids:
                    continue
                cross_comment_ids.add(identifier)
                quarter = quarter_for(comment, line_number, filtered_path)
                if int(quarter[:4]) >= STATISTICS_START_YEAR:
                    if last_quarter is None or quarter > last_quarter:
                        last_quarter = quarter
                    for group in groups:
                        if author:
                            distinct_users[quarter][group].add(author)
                    overlap = overlap_group_name(groups)
                    if overlap:
                        overlap_comment_counts[quarter][overlap] += 1
                        overlap_distinct_users[quarter][overlap].add(author)
                for group in GROUPS:
                    if group not in groups:
                        continue
                    group_outputs[group].write(line)
                    cross_group_counts[group] += 1
                    if int(quarter[:4]) >= STATISTICS_START_YEAR:
                        counts[quarter][group] += 1

        for group, path in group_spills.items():
            path.replace(output_dir / f"{group}.jsonl")
        baseline_spill.replace(output_dir / "baseline.jsonl")
    finally:
        baseline_spill.unlink(missing_ok=True)
        for path in group_spills.values():
            path.unlink(missing_ok=True)

    quarters = (
        quarter_sequence(f"{STATISTICS_START_YEAR}-Q1", last_quarter)
        if last_quarter and int(last_quarter[:4]) >= STATISTICS_START_YEAR
        else []
    )
    total_seconds = time.monotonic() - run_started
    run_info = [
        ("run started", run_started_at),
        ("baseline input", str(baseline_path)),
        ("grouped input", str(grouped_dir)),
        ("age input", str(age_dir)),
        ("filtered input", str(filtered_path)),
        ("relevant authors", f"{len(groups_by_author):,}"),
        ("original comment IDs", f"{len(original_ids):,}"),
        ("baseline comments processed", f"{baseline_comments:,}"),
        ("baseline comments moved", f"{baseline_removed:,}"),
        ("filtered comments processed", f"{filtered_comments:,}"),
        ("new unique comments", f"{len(cross_comment_ids):,}"),
        ("total runtime", f"{total_seconds:.2f} seconds"),
    ]
    write_statistics(
        output_dir / "quarterly_statistics.xlsx",
        quarters,
        counts,
        distinct_users,
        overlap_comment_counts,
        overlap_distinct_users,
        run_info,
    )

    for group in GROUPS:
        print(f"{group}.jsonl: {cross_group_counts[group]:,} additional comments")
    print(f"baseline.jsonl: {baseline_comments - baseline_removed:,} remaining comments")
    print(f"Quarterly statistics -> {output_dir / 'quarterly_statistics.xlsx'}")
    with (output_dir / "cross_subreddit_runtime.log").open("a", encoding="utf-8") as log_file:
        log_file.write(
            f"{datetime.now().isoformat(timespec='seconds')} | "
            f"authors={len(groups_by_author)} | baseline={baseline_comments} | "
            f"filtered={filtered_comments} | new_comments={len(cross_comment_ids)} | "
            f"total={total_seconds:.2f}s\n"
        )


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--baseline", type=Path, default=DEFAULT_BASELINE)
    parser.add_argument("--grouped-dir", type=Path, default=GROUPED_DIR)
    parser.add_argument("--age-dir", type=Path, default=AGE_DIR)
    parser.add_argument("--filtered", type=Path, default=DEFAULT_FILTERED_INPUT)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    args = parser.parse_args()
    process_cross_subreddit_data(
        args.baseline,
        args.grouped_dir,
        args.filtered,
        args.output_dir,
        args.age_dir,
        load_deleted_authors(GROUPED_FLAIRS_PATH),
    )


if __name__ == "__main__":
    main()