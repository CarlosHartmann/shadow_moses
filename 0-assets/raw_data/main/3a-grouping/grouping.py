"""Split filtered Reddit comments into demographic and political groups."""

from __future__ import annotations

import argparse
import json
import time
from collections import defaultdict
from contextlib import ExitStack
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterator

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE


HERE = Path(__file__).resolve().parent
DEFAULT_INPUT = HERE.parent / "2-filtering" / "filtered_comments.jsonl"
DEFAULT_OUTPUT_DIR = HERE / "grouped_subreddit_data"
SUBREDDIT_LISTS_DIR = HERE / "subreddit_lists"
STATISTICS_START_YEAR = 2010
EXCEL_MAX_CELL_CHARS = 32_767
PROGRESS_EVERY = 1_000_000
RUNTIME_LOG_NAME = "grouping_runtime.log"
GROUPED_FLAIRS_NAME = "grouped_flairs.xlsx"
FLAIR_COLUMNS = ["subreddit", "flair", "relevant"]
AGE_GROUP = "age"
# In the age sheet a flair is relevant if any of these columns holds a value.
AGE_ANNOTATION_COLUMNS = (
    "given_age",
    "given_birthyear",
    "given_agerange",
    "given_birthyear_range",
    "given_generation",
    "given_description",
)
LEGACY_RELEVANT_MARKER = "X"
UNKNOWN_AUTHORS = {"", "[deleted]", "[removed]"}
USER_OVERRIDE_SHEET = "user_group_overrides"
# Users with an override keep only the decided group of the family ("none" keeps neither).
# A decided group outside the family (e.g. gender_political_progressive) also takes the user's
# comments from the family's subreddits regardless of flair and score.
OVERRIDE_FAMILIES = {
    "age": (AGE_GROUP,),
    "gender": ("gender_male", "gender_female"),
    "political": ("political_conservative", "political_progressive"),
    "gender_political": ("gender_political_conservative", "gender_political_progressive"),
}
GROUP_ORDER = (
    AGE_GROUP,
    "gender_male",
    "gender_female",
    "political_conservative",
    "political_progressive",
    "gender_political_conservative",
    "gender_political_progressive",
)
# Every group follows grouped_flairs.xlsx (see flair_admits for comments before 2012).
FLAIR_RELEVANCE_GROUPS = frozenset(
    {
        AGE_GROUP,
        "gender_male",
        "gender_female",
        "political_conservative",
        "political_progressive",
        "gender_political_conservative",
        "gender_political_progressive",
    }
)


def read_subreddit_list(path: Path) -> set[str]:
    """One subreddit per line, case-insensitive; blank lines and # comments are skipped."""
    if not path.exists():
        return set()
    return {
        line.strip().casefold()
        for line in path.read_text(encoding="utf-8").splitlines()
        if line.strip() and not line.lstrip().startswith("#")
    }


def load_group_subreddits(directory: Path) -> dict[str, set[str]]:
    """Read the subreddits of each group from <group>.txt files."""
    return {name: read_subreddit_list(directory / f"{name}.txt") for name in GROUP_ORDER}


GROUP_SUBREDDITS = load_group_subreddits(SUBREDDIT_LISTS_DIR)
# Before 2012, comments from these subreddits count for the group without a relevant flair
# (score of 5 or more). A group without a list (e.g. age) only takes relevant flairs then.
HIGH_LIKELIHOOD_SUBREDDITS = load_group_subreddits(SUBREDDIT_LISTS_DIR / "high_likelihood")
for _group, _subreddits in HIGH_LIKELIHOOD_SUBREDDITS.items():
    if not _subreddits <= GROUP_SUBREDDITS[_group]:
        raise ValueError(
            f"high_likelihood/{_group}.txt lists subreddits missing from {_group}.txt: "
            + ", ".join(sorted(_subreddits - GROUP_SUBREDDITS[_group]))
        )
SCORE_FILTERED_GROUPS = {
    "political_conservative",
    "political_progressive",
    "gender_political_conservative",
    "gender_political_progressive",
}
FLAIR_START = datetime(2012, 1, 1, tzinfo=timezone.utc).timestamp()
ALL_TARGET_SUBREDDITS = set().union(*GROUP_SUBREDDITS.values())
GROUP_NAMES = tuple(GROUP_SUBREDDITS)
# Age statistics come from the 3b group_by scripts, once the age data is split.
STATISTICS_GROUP_NAMES = tuple(group for group in GROUP_NAMES if group != AGE_GROUP)


def iter_comments(path: Path) -> Iterator[tuple[int, str, dict]]:
    with path.open("r", encoding="utf-8") as input_file:
        for line_number, line in enumerate(input_file, start=1):
            if not line.strip():
                continue
            try:
                comment = json.loads(line)
            except json.JSONDecodeError as error:
                raise ValueError(f"Invalid JSON on line {line_number}: {error}") from error
            if not isinstance(comment, dict):
                raise ValueError(f"Expected a JSON object on line {line_number}.")
            yield line_number, line if line.endswith("\n") else line + "\n", comment


def quarter_for(comment: dict, line_number: int) -> str:
    try:
        timestamp = float(comment["created_utc"])
        created_at = datetime.fromtimestamp(timestamp, tz=timezone.utc)
    except (KeyError, TypeError, ValueError, OverflowError, OSError) as error:
        raise ValueError(f"Missing or invalid created_utc on line {line_number}.") from error
    quarter = (created_at.month - 1) // 3 + 1
    return f"{created_at.year}-Q{quarter}"


def score_qualifies(comment: dict) -> bool:
    try:
        return float(comment.get("score", 0)) >= 5
    except (TypeError, ValueError):
        return False


def matching_groups(comment: dict) -> list[str]:
    subreddit = str(comment.get("subreddit", "")).casefold()
    groups = [
        group
        for group, subreddits in GROUP_SUBREDDITS.items()
        if subreddit in subreddits
    ]
    qualifies = score_qualifies(comment)
    return [
        group
        for group in groups
        if group not in SCORE_FILTERED_GROUPS or qualifies
    ]


def flair_admits(
    group: str,
    comment: dict,
    pair: tuple[str, str],
    relevant_flairs: dict[str, set[tuple[str, str]]],
) -> bool:
    """Whether a comment of the group's subreddits counts for the group (score filter aside).

    From 2012 on, only a relevant flair admits it. Before 2012 a relevant flair admits it as
    well (flairs seen on older comments were set later but are still informative), and so does
    a high-likelihood subreddit of the group if the comment has a score of 5 or more.
    """
    if pair[1] and pair in relevant_flairs[group]:
        return True
    return (
        float(comment["created_utc"]) < FLAIR_START
        and pair[0] in HIGH_LIKELIHOOD_SUBREDDITS[group]
        and score_qualifies(comment)
    )


def excel_value(value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)[:EXCEL_MAX_CELL_CHARS]
    return value


def quarter_sequence(first: str, last: str) -> list[str]:
    year, quarter = map(int, first.replace("-Q", " ").split())
    last_year, last_quarter = map(int, last.replace("-Q", " ").split())
    quarters = []
    while (year, quarter) <= (last_year, last_quarter):
        quarters.append(f"{year}-Q{quarter}")
        year, quarter = (year + 1, 1) if quarter == 4 else (year, quarter + 1)
    return quarters


def write_statistics(
    output_path: Path, quarters: list[str], counts: dict, run_info: list[tuple[str, str]]
) -> None:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Quarterly counts")
    sheet.append(["quarter", *STATISTICS_GROUP_NAMES])
    for quarter in quarters:
        sheet.append([quarter, *(counts[quarter][group] for group in STATISTICS_GROUP_NAMES)])
    info_sheet = workbook.create_sheet("Run info")
    for row in run_info:
        info_sheet.append(list(row))
    workbook.save(output_path)


def flair_of(comment: dict) -> str:
    return excel_value(str(comment.get("author_flair_text") or "").strip())


def update_grouped_flairs(
    path: Path, combos: dict[str, set[tuple[str, str]]]
) -> dict[str, int]:
    """Create grouped_flairs.xlsx or append unseen combinations, keeping existing 'relevant' values."""
    if path.exists():
        workbook = load_workbook(path)
    else:
        workbook = Workbook()
        del workbook[workbook.sheetnames[0]]

    added = {}
    for group in GROUP_NAMES:
        if group in workbook.sheetnames:
            sheet = workbook[group]
        else:
            sheet = workbook.create_sheet(group)
            sheet.append(
                ["subreddit", "flair", *AGE_ANNOTATION_COLUMNS]
                if group == AGE_GROUP
                else FLAIR_COLUMNS
            )
        known = {
            (str(row[0] or ""), str(row[1] or ""))
            for row in sheet.iter_rows(min_row=2, max_col=2, values_only=True)
        }
        new_combos = sorted(combos[group] - known)
        for subreddit, flair in new_combos:
            # Unseen flairs stay blank, i.e. irrelevant until annotated.
            sheet.append([subreddit, flair])
        added[group] = len(new_combos)
    workbook.save(path)
    return added


def load_relevant_flairs(
    path: Path,
) -> tuple[dict[str, set[tuple[str, str]]], dict[str, set[tuple[str, str]]]] | None:
    """Return known and relevant pairs, or None when no annotation workbook exists."""
    if not path.exists():
        return None
    workbook = load_workbook(path, read_only=True)
    known: dict[str, set[tuple[str, str]]] = {group: set() for group in GROUP_NAMES}
    relevant: dict[str, set[tuple[str, str]]] = {group: set() for group in GROUP_NAMES}
    for group in GROUP_NAMES:
        if group not in workbook.sheetnames:
            continue
        if group == AGE_GROUP:
            rows = workbook[group].iter_rows(values_only=True)
            header = [str(name or "") for name in next(rows, ())]
            columns = [header.index(name) for name in AGE_ANNOTATION_COLUMNS if name in header]
            for row in rows:
                row = (*row, *[None] * (len(header) - len(row)))
                pair = (str(row[0] or ""), str(row[1] or ""))
                known[group].add(pair)
                # A flair can occupy several rows; any annotated row makes it relevant.
                # In given_generation, "X" is Generation X, not the legacy marker.
                if any(
                    str(row[index] or "").strip()
                    and (
                        header[index] == "given_generation"
                        or str(row[index]).strip().upper() != LEGACY_RELEVANT_MARKER
                    )
                    for index in columns
                ):
                    relevant[group].add(pair)
            continue
        for row in workbook[group].iter_rows(min_row=2, max_col=3, values_only=True):
            pair = (str(row[0] or ""), str(row[1] or ""))
            known[group].add(pair)
            if str(row[2] or "").strip():
                relevant[group].add(pair)
    workbook.close()
    return known, relevant


def load_user_overrides(
    path: Path,
) -> tuple[dict[str, set[str]], dict[str, dict[str, tuple[str, ...]]]]:
    """Per author: groups they are barred from, and forced groups with the family they replace."""
    exclusions: dict[str, set[str]] = defaultdict(set)
    inclusions: dict[str, dict[str, tuple[str, ...]]] = defaultdict(dict)
    if not path.exists():
        return exclusions, inclusions
    workbook = load_workbook(path, read_only=True)
    if USER_OVERRIDE_SHEET in workbook.sheetnames:
        for user, family, keep in workbook[USER_OVERRIDE_SHEET].iter_rows(
            min_row=2, max_col=3, values_only=True
        ):
            if user and family in OVERRIDE_FAMILIES:
                author = str(user).casefold()
                exclusions[author] |= set(OVERRIDE_FAMILIES[family]) - {keep}
                if keep in GROUP_ORDER and keep not in OVERRIDE_FAMILIES[family]:
                    inclusions[author][keep] = OVERRIDE_FAMILIES[family]
    workbook.close()
    return exclusions, inclusions


def format_duration(seconds: float) -> str:
    hours, remainder = divmod(int(seconds), 3600)
    minutes, secs = divmod(remainder, 60)
    return f"{hours}h {minutes:02d}m {secs:02d}s"


def group_comments(input_path: Path, output_dir: Path) -> None:
    run_started = time.monotonic()
    run_started_at = datetime.now().isoformat(timespec="seconds")
    output_dir.mkdir(parents=True, exist_ok=True)
    baseline_path = output_dir / "baseline.jsonl"
    spill_paths = {group: output_dir / f".{group}.spill.jsonl" for group in GROUP_NAMES}
    group_counts = defaultdict(int)
    counts = defaultdict(lambda: defaultdict(int))
    flair_combos: dict[str, set[tuple[str, str]]] = {group: set() for group in GROUP_NAMES}
    last_quarter = None
    total_comments = 0

    print(f"Reading {input_path}", flush=True)
    user_exclusions, user_inclusions = load_user_overrides(HERE / GROUPED_FLAIRS_NAME)
    flair_annotations = load_relevant_flairs(HERE / GROUPED_FLAIRS_NAME)
    if flair_annotations is None:
        print(f"No {GROUPED_FLAIRS_NAME} yet; flair relevance is not applied.", flush=True)
        relevant_flairs = None
    else:
        _, relevant_flairs = flair_annotations
    try:
        with ExitStack() as stack:
            baseline_file = stack.enter_context(baseline_path.open("w", encoding="utf-8"))
            spill_files = {
                group: stack.enter_context(path.open("w", encoding="utf-8"))
                for group, path in spill_paths.items()
            }
            for line_number, line, comment in iter_comments(input_path):
                total_comments += 1
                quarter = quarter_for(comment, line_number)
                if last_quarter is None or quarter > last_quarter:
                    last_quarter = quarter
                author = str(comment.get("author") or "").casefold()
                subreddit = str(comment.get("subreddit", "")).casefold()
                groups = matching_groups(comment)
                # Ignoring flair and score, but only for the user's comments from the subreddits of
                # the replaced family, i.e. where they would have been in the replaced groups.
                forced_groups = {
                    forced_group
                    for forced_group, replaced in user_inclusions.get(author, {}).items()
                    if any(subreddit in GROUP_SUBREDDITS[name] for name in replaced)
                }
                groups.extend(sorted(forced_groups - set(groups)))
                for group in groups:
                    pair = (subreddit, flair_of(comment))
                    if subreddit in GROUP_SUBREDDITS[group]:
                        flair_combos[group].add(pair)
                    # Comments without a known author cannot be balanced per user.
                    if author in UNKNOWN_AUTHORS or group in user_exclusions.get(author, ()):
                        continue
                    forced = group in forced_groups
                    if (
                        not forced
                        and relevant_flairs is not None
                        and group in FLAIR_RELEVANCE_GROUPS
                        and not flair_admits(group, comment, pair, relevant_flairs)
                    ):
                        continue
                    spill_files[group].write(line)
                    group_counts[group] += 1
                    if int(quarter[:4]) >= STATISTICS_START_YEAR:
                        counts[quarter][group] += 1
                subreddit = str(comment.get("subreddit", "")).casefold()
                if subreddit not in ALL_TARGET_SUBREDDITS:
                    baseline_file.write(line)
                if total_comments % PROGRESS_EVERY == 0:
                    elapsed = time.monotonic() - run_started
                    print(
                        f"  {total_comments:,} comments, {format_duration(elapsed)} elapsed, "
                        f"{total_comments / elapsed:,.0f} comments/s",
                        flush=True,
                    )
        pass_seconds = time.monotonic() - run_started
        print(
            f"Pass over input finished: {total_comments:,} comments in "
            f"{format_duration(pass_seconds)}",
            flush=True,
        )

        quarters = (
            quarter_sequence(f"{STATISTICS_START_YEAR}-Q1", last_quarter)
            if last_quarter and int(last_quarter[:4]) >= STATISTICS_START_YEAR
            else []
        )
        flairs_path = HERE / GROUPED_FLAIRS_NAME
        flairs_existed = flairs_path.exists()
        added = update_grouped_flairs(flairs_path, flair_combos)
        print(
            f"{GROUPED_FLAIRS_NAME} {'updated' if flairs_existed else 'created'}: "
            + ", ".join(f"{group}={count:,} new" for group, count in added.items()),
            flush=True,
        )
        statistics_path = output_dir / "quarterly_statistics.xlsx"

        output_started = time.monotonic()
        for group in GROUP_NAMES:
            output_path = output_dir / f"{group}.jsonl"
            spill_paths[group].replace(output_path)
            print(
                f"  {group}.jsonl: {group_counts[group]:,} comments",
                flush=True,
            )
        output_seconds = time.monotonic() - output_started
    finally:
        for path in spill_paths.values():
            path.unlink(missing_ok=True)

    print(f"Baseline comments -> {baseline_path}")
    print(f"Quarterly statistics -> {statistics_path}")
    print(f"Group JSONL files -> {output_dir}")

    total_seconds = time.monotonic() - run_started
    run_info = [
        ("run started", run_started_at),
        ("input", str(input_path)),
        ("comments processed", f"{total_comments:,}"),
        ("input pass", format_duration(pass_seconds)),
        ("group output writing", format_duration(output_seconds)),
        ("total runtime", format_duration(total_seconds)),
    ]
    write_statistics(statistics_path, quarters, counts, run_info)
    print(f"Total runtime: {format_duration(total_seconds)}")
    with (output_dir / RUNTIME_LOG_NAME).open("a", encoding="utf-8") as log_file:
        log_file.write(
            f"{datetime.now().isoformat(timespec='seconds')} | input={input_path} | "
            f"comments={total_comments} | pass={format_duration(pass_seconds)} | "
            f"group_output={format_duration(output_seconds)} | "
            f"total={format_duration(total_seconds)}\n"
        )


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    args = parser.parse_args()
    group_comments(args.input, args.output_dir)


if __name__ == "__main__":
    main()