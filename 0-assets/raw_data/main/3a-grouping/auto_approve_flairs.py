import re
from pathlib import Path

import openpyxl

PATH = Path(__file__).parent / "grouped_flairs.xlsx"
SHEET = "gender_female_or_diverse"
PATTERNS = [
    "amab", "afab", "mtf", "ftm", "hrt", "femme", "butch", "genderfluid",
    "agender", "she/her", "m2f", "f2m", "pre everything", "pre-everything",
    "nb", "enby", "non-binary", "nonbinary", "trans", "trans-",
]
REGEX_PATTERNS = [
    r"\d\d/? ?F", r"y/o girl", r"y/o woman", r"they/them", r"E since", r"T since",
    r"trans ?woman", r"trans ?girl", r"trans ?guy", r"trans ?man", r"transgender",
    r"cis woman", r"cis ?girl", r"genderqueer", r"demi-boy", r"demi-girl",
]

wb = openpyxl.load_workbook(PATH)
ws = wb[SHEET]
header = [c.value for c in ws[1]]
flair_col = header.index("flair") + 1
auto_col = header.index("automatically_approved") + 1

regex = re.compile(
    "|".join(
        [rf"\b(?:{re.escape(p)})\b" for p in PATTERNS]
        + [rf"\b(?:{p})\b" for p in REGEX_PATTERNS]
    ),
    re.IGNORECASE,
)

count = 0
for row in range(2, ws.max_row + 1):
    flair = ws.cell(row, flair_col).value
    if flair and regex.search(str(flair)):
        ws.cell(row, auto_col).value = "X"
        count += 1

wb.save(PATH)
print(f"Marked {count} rows")
