"""Write the decisions in overlap_audit.xlsx into grouped_flairs.xlsx as per-user override sheets.

- age_user_overrides: "birth year(s)" rows of the Decisions sheet; read by 3b-age-grouping/deduce_birthyear.py.
- user_group_overrides: political and age-"delete" rows of the Decisions sheet and the Gender report sheet;
  read by 3a-grouping/grouping.py. Both sheets are rebuilt on every run.
"""

from __future__ import annotations

import argparse
import re
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
GROUPING_DIR = HERE.parent / "3a-grouping"
AUDIT_PATH = HERE / "overlap_audit.xlsx"
FLAIRS_PATH = GROUPING_DIR / "grouped_flairs.xlsx"
BACKUP_DIR = GROUPING_DIR / "annotation_backup"
AGE_SHEET = "age_user_overrides"
GROUP_SHEET = "user_group_overrides"
BIRTH_YEAR = re.compile(r"^birth years? (\d{4})(?:-(\d{4}))?$")
# Gender report decision -> group to keep; a group outside male/female also takes the user's
# comments from its subreddits ("none" removes the user from both gender groups).
GENDER_DECISIONS = {
    "male": "gender_male",
    "female": "gender_female",
    "diverse": "gender_political_progressive",
    "delete": "none",
}
POLITICAL_FAMILY = "political"
GENDER_POLITICAL_FAMILY = "gender_political"
AGE_FAMILY = "age"
GENDER_FAMILY = "gender"


def collect_decisions(audit_path: Path):
    workbook = load_workbook(audit_path, read_only=True)
    age: dict[str, tuple[int, int]] = {}
    groups: dict[tuple[str, str], str] = {}
    for user, value, _ in workbook["Decisions"].iter_rows(min_row=2, max_col=3, values_only=True):
        value = str(value or "").strip()
        match = BIRTH_YEAR.match(value)
        if user and match:
            age[str(user).casefold()] = (int(match[1]), int(match[2] or match[1]))
        elif user and (value.startswith("political_") or value == "none"):
            groups[(str(user).casefold(), POLITICAL_FAMILY)] = value
        elif user and value.startswith("gender_political_"):
            groups[(str(user).casefold(), GENDER_POLITICAL_FAMILY)] = value
        elif user and value == "delete":
            groups[(str(user).casefold(), AGE_FAMILY)] = "none"
    for row in workbook["Gender report"].iter_rows(min_row=2, max_col=2, values_only=True):
        decision, user = (str(cell or "").strip() for cell in row)
        if user and decision in GENDER_DECISIONS:
            groups[(user.casefold(), GENDER_FAMILY)] = GENDER_DECISIONS[decision]
    workbook.close()
    return age, groups


def write_sheet(workbook, name: str, header: list[str], rows: list[list]) -> None:
    if name in workbook.sheetnames:
        del workbook[name]
    sheet = workbook.create_sheet(name)
    sheet.append(header)
    for row in rows:
        sheet.append(row)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--audit", type=Path, default=AUDIT_PATH)
    parser.add_argument("--flairs", type=Path, default=FLAIRS_PATH)
    args = parser.parse_args()

    age, groups = collect_decisions(args.audit)
    BACKUP_DIR.mkdir(exist_ok=True)
    stamp = datetime.now().strftime("%Y-%m-%d_%Hh-%Mm-%Ss")
    shutil.copy(args.flairs, BACKUP_DIR / f"grouped_flairs_backup_{stamp}.xlsx")
    workbook = load_workbook(args.flairs)
    write_sheet(
        workbook,
        AGE_SHEET,
        ["user", "birthyear_lo", "birthyear_hi"],
        [[user, lo, hi] for user, (lo, hi) in sorted(age.items())],
    )
    write_sheet(
        workbook,
        GROUP_SHEET,
        ["user", "family", "keep_group"],
        [[user, family, keep] for (user, family), keep in sorted(groups.items())],
    )
    workbook.save(args.flairs)
    print(f"{AGE_SHEET}: {len(age)} users; {GROUP_SHEET}: {len(groups)} entries")


if __name__ == "__main__":
    main()
