"""Explain why users fall into mutually exclusive groups.

Membership follows cross_subreddit.py (original group comments from 2012 on). For every
user in a contradictory combination, the evidence rows list each (group, subreddit, flair,
deduced birth year range) that placed them there.
"""

from __future__ import annotations

import argparse
import sys
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "4-cross-subreddit-data"))

from cross_subreddit import (  # noqa: E402
    AGE_DIR,
    AUTHOR_SELECTION_START,
    GROUPED_DIR,
    GROUPS,
    author_key,
    group_path,
    iter_comments,
)

GENERATIONS = {"age_silent", "age_boomer", "age_X", "age_millennial", "age_Z"}
OLD_GENERATIONS = {"age_silent", "age_boomer", "age_X"}  # born 1980 or earlier
YOUNG_GENERATIONS = {"age_millennial", "age_Z"}  # born 1981 or later


def contradictions(groups: set[str]) -> dict[str, set[str]]:
    """Map each violated rule to the groups involved."""
    found: dict[str, set[str]] = {}

    def add(rule: str, involved: set[str]) -> None:
        found[rule] = involved

    if len(groups & GENERATIONS) >= 2:
        add("different generations", groups & GENERATIONS)
    if {"age_old", "age_young"} <= groups:
        add("old and young", {"age_old", "age_young"})
    if "age_old" in groups and groups & YOUNG_GENERATIONS:
        add("old and young generation", {"age_old"} | groups & YOUNG_GENERATIONS)
    if "age_young" in groups and groups & OLD_GENERATIONS:
        add("young and old generation", {"age_young"} | groups & OLD_GENERATIONS)
    if {"gender_male", "gender_female_or_diverse"} <= groups:
        add("male and female_or_diverse", {"gender_male", "gender_female_or_diverse"})
    if {"political_conservative", "political_progressive"} <= groups:
        add("conservative and progressive", {"political_conservative", "political_progressive"})
    if {"gender_political_conservative", "gender_political_progressive"} <= groups:
        add(
            "gender-conservative and gender-progressive",
            {"gender_political_conservative", "gender_political_progressive"},
        )
    return found


def iso_date(comment: dict) -> str:
    return datetime.fromtimestamp(float(comment["created_utc"]), tz=timezone.utc).date().isoformat()


def run(grouped_dir: Path, age_dir: Path, output_path: Path) -> None:
    groups_by_author: dict[str, set[str]] = defaultdict(set)
    for group in GROUPS:
        path = group_path(group, grouped_dir, age_dir)
        for _, _, comment in iter_comments(path):
            author = author_key(comment)
            if author and float(comment["created_utc"]) >= AUTHOR_SELECTION_START:
                groups_by_author[author].add(group)

    violations = {
        author: found
        for author, groups in groups_by_author.items()
        if (found := contradictions(groups))
    }

    # (author, group, subreddit, flair text, css class, birth range) -> [count, first, last, example id]
    evidence: dict[tuple, list] = {}
    for group in GROUPS:
        path = group_path(group, grouped_dir, age_dir)
        for _, _, comment in iter_comments(path):
            author = author_key(comment)
            if author not in violations or float(comment["created_utc"]) < AUTHOR_SELECTION_START:
                continue
            birth_range = comment.get("deduced_birthyear_range")
            key = (
                author,
                group,
                comment.get("subreddit"),
                comment.get("author_flair_text"),
                comment.get("author_flair_css_class"),
                "-".join(map(str, birth_range)) if birth_range else None,
            )
            date = iso_date(comment)
            row = evidence.get(key)
            if row is None:
                evidence[key] = [1, date, date, comment.get("id")]
            else:
                row[0] += 1
                row[1] = min(row[1], date)
                row[2] = max(row[2], date)

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["rule", "users"])
    rule_counts: dict[str, int] = defaultdict(int)
    for found in violations.values():
        for rule in found:
            rule_counts[rule] += 1
    for rule, count in sorted(rule_counts.items()):
        summary.append([rule, count])

    users = workbook.create_sheet("Users")
    users.append(["user", "rule", "groups involved", "all groups"])
    for author in sorted(violations):
        for rule, involved in violations[author].items():
            users.append([
                author,
                rule,
                " + ".join(g for g in GROUPS if g in involved),
                " + ".join(g for g in GROUPS if g in groups_by_author[author]),
            ])

    sheet = workbook.create_sheet("Evidence")
    sheet.append([
        "user", "rules", "group", "subreddit", "flair text", "flair css class",
        "deduced birth years", "comments", "first date", "last date", "example comment id",
    ])
    for key in sorted(evidence, key=lambda k: (k[0], GROUPS.index(k[1]), str(k[2]))):
        involved_groups = set().union(*violations[key[0]].values())
        if key[1] not in involved_groups:
            continue
        count, first, last, example = evidence[key]
        sheet.append([
            key[0], "; ".join(violations[key[0]]), *key[1:], count, first, last, example,
        ])
    workbook.save(output_path)

    for rule, count in sorted(rule_counts.items()):
        print(f"{rule}: {count:,} users")
    print(f"Audit -> {output_path}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--grouped-dir", type=Path, default=GROUPED_DIR)
    parser.add_argument("--age-dir", type=Path, default=AGE_DIR)
    parser.add_argument("--output", type=Path, default=HERE / "overlap_audit.xlsx")
    args = parser.parse_args()
    run(args.grouped_dir, args.age_dir, args.output)


if __name__ == "__main__":
    main()
