"""Derive the rule-based age decisions of the overlap audit from the flair-derived birth years.

A user needs an age decision if their age comments from 2012 on (3a age.jsonl, flair-derived
birth year ranges only, no overrides) fall into different generations or into both young and old.
For each such user, in this order:
1. Intersection of all flair-derived ranges (highest lo, lowest hi). If the span crosses a group
   boundary, the majority bucket of the 3b scripts decides later.
2. Otherwise, the intersection of the range at each flair's first appearance (stale flairs ignored).
3. Otherwise there is no common year. Decisions made by hand (and deletions) are kept; users whose
   earlier decision was the average of the midpoints of their distinct ranges get that average
   recomputed; anyone else is marked "needs decision", with the average as a suggestion.

The age rows of the Decisions sheet in overlap_audit.xlsx are rebuilt from this (hand-made rows are
kept as they are) and the Age report sheet is rewritten. A backup of the workbook is made first.
Run apply_decisions.py afterwards to pass the decisions on to grouped_flairs.xlsx.
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
AGE_DIR = HERE.parent / "3b-age-grouping"
sys.path.insert(0, str(AGE_DIR))

from age_split import majority_bucket  # noqa: E402
from deduce_birthyear import (  # noqa: E402
    DEFAULT_INPUT,
    GROUPED_FLAIRS_PATH,
    birth_range_for,
    flair_of,
    flair_year,
    load_annotations,
)
from group_by_age import BUCKETS as YOUNG_OLD  # noqa: E402
from group_by_generation import GENERATIONS  # noqa: E402

AUDIT_PATH = HERE / "overlap_audit.xlsx"
BACKUP_DIR = HERE / "backup"
SELECTION_START = datetime(2012, 1, 1, tzinfo=timezone.utc).timestamp()
ANNOTATION_COLUMNS = (
    "given_age",
    "given_birthyear",
    "given_generation",
    "given_agerange",
    "given_birthyear_range",
    "given_description",
)
BIRTH_YEAR = re.compile(r"^birth years? (\d{4})(?:-(\d{4}))?$")
NEEDS_DECISION = "needs decision"
REASON_INTERSECTION = "Intersection of all flair-derived ranges (highest lo, lowest hi)"
REASON_CROSSING = (
    "Intersection of all flair-derived ranges; span crosses a group boundary, majority bucket decides"
)
REASON_FIRST = "Intersection of the range from each flair's first appearance (stale flairs ignored)"
REASON_AVERAGE = "Average of the midpoints of all flair-derived ranges (young users, born 1995-1999)"
RULE_REASONS = (REASON_INTERSECTION, REASON_CROSSING, REASON_FIRST, REASON_AVERAGE)
REPORT_HEADER = [
    "user",
    "status",
    "common birth years (max lo - min hi)",
    "subreddit",
    "flair",
    "years posted (comments)",
    *ANNOTATION_COLUMNS,
    "deduced birth year ranges",
]


def comment_year(comment: dict) -> int:
    return datetime.fromtimestamp(float(comment["created_utc"]), tz=timezone.utc).year


def load_flair_ranges(input_path: Path, annotations: dict) -> dict:
    """user -> (subreddit, flair) -> comment year -> [birth range, comment count]."""
    users: dict = defaultdict(lambda: defaultdict(dict))
    with input_path.open(encoding="utf-8") as source:
        for line in source:
            if not line.strip():
                continue
            comment = json.loads(line)
            if float(comment["created_utc"]) < SELECTION_START:
                continue
            key = (str(comment.get("subreddit", "")).casefold(), flair_of(comment))
            annotation = annotations.get(key)
            birth_range = birth_range_for(annotation, flair_year(comment)) if annotation else None
            if not birth_range:
                continue
            years = users[str(comment.get("author") or "").casefold()][key]
            entry = years.setdefault(comment_year(comment), [birth_range, 0])
            entry[1] += 1
    return users


def all_ranges(flairs: dict) -> list[tuple[int, int]]:
    return [birth_range for years in flairs.values() for birth_range, _ in years.values()]


def is_contradictory(flairs: dict) -> bool:
    ranges = all_ranges(flairs)
    generations = {majority_bucket(r, GENERATIONS) for r in ranges} - {None}
    young_old = {majority_bucket(r, YOUNG_OLD) for r in ranges} - {None}
    return len(generations) > 1 or len(young_old) > 1


def crosses_boundary(low: int, high: int) -> bool:
    return len({majority_bucket((year, year), GENERATIONS) for year in range(low, high + 1)}) > 1


def intersection(ranges: list[tuple[int, int]]) -> tuple[int, int]:
    return max(low for low, _ in ranges), min(high for _, high in ranges)


def midpoint_average(flairs: dict) -> int:
    midpoints = [(low + high) / 2 for low, high in set(all_ranges(flairs))]
    return round(sum(midpoints) / len(midpoints))


def derive(flairs: dict) -> tuple[str, tuple[int, int] | None, str]:
    """(kind, birth year range or None, common span as shown in the Age report)."""
    low, high = intersection(all_ranges(flairs))
    if low <= high:
        kind = "crossing" if crosses_boundary(low, high) else "intersection"
        return kind, (low, high), f"{low}-{high}"
    first = [years[min(years)][0] for years in flairs.values()]
    low, high = intersection(first)
    if low <= high:
        return "first", (low, high), f"{low}-{high}"
    return "none", None, f"{low}-{high} (first appearance)"


def birth_year_value(birth_range: tuple[int, int]) -> str:
    low, high = birth_range
    return f"birth year {low}" if low == high else f"birth years {low}-{high}"


def years_posted(years: dict) -> str:
    return ", ".join(f"{year} ({years[year][1]})" for year in sorted(years))


def ranges_text(years: dict) -> str:
    return "; ".join(f"{low}-{high}" for low, high in sorted({r for r, _ in years.values()}))


def flairs_text(flairs: dict) -> str:
    return "\n".join(
        f"{subreddit} | {flair} | {years_posted(years)} | {ranges_text(years)}"
        for (subreddit, flair), years in sorted(flairs.items())
    )


def ranges_summary(flairs: dict) -> str:
    low, high = intersection(all_ranges(flairs))
    return f"{low} / {high}"


def run(input_path: Path, flairs_path: Path, audit_path: Path) -> None:
    annotations, _ = load_annotations(flairs_path)
    users = load_flair_ranges(input_path, annotations)
    contradictory = {user for user, flairs in users.items() if is_contradictory(flairs)}

    BACKUP_DIR.mkdir(exist_ok=True)
    stamp = datetime.now().strftime("%Y-%m-%d_%Hh-%Mm-%Ss")
    shutil.copy(audit_path, BACKUP_DIR / f"overlap_audit_backup_{stamp}.xlsx")
    workbook = load_workbook(audit_path)
    decisions = workbook["Decisions"]
    rows = [list(row) for row in decisions.iter_rows(values_only=True)]
    header, rows = rows[0][:5], [row[:5] + [None] * (5 - len(row[:5])) for row in rows[1:]]

    kept, earlier = [], {}
    for row in rows:
        user, value, reason = str(row[0] or "").casefold(), str(row[1] or ""), str(row[2] or "")
        if (BIRTH_YEAR.match(value) or value == NEEDS_DECISION) and (
            reason in RULE_REASONS or value == NEEDS_DECISION
        ):
            earlier[user] = reason
        else:
            kept.append(row)
    decided_by_hand = {
        str(row[0]).casefold(): row for row in kept if BIRTH_YEAR.match(str(row[1] or ""))
    }
    deleted = {str(row[0]).casefold() for row in kept if row[1] == "delete"}

    new_rows, status_of, outcome = [], {}, Counter()
    hand_differs = []
    for user in sorted(contradictory):
        flairs = users[user]
        kind, birth_range, span = derive(flairs)
        if user in decided_by_hand or user in deleted:
            outcome["kept (decided by hand)" if user in decided_by_hand else "kept (delete)"] += 1
            status_of[user] = ("resolved" if birth_range else NEEDS_DECISION, span)
            hand = decided_by_hand.get(user)
            # Only worth a look where the rules now find a common year that differs.
            if hand is not None and birth_range and birth_year_value(birth_range) != hand[1]:
                hand_differs.append((user, hand[1], birth_year_value(birth_range)))
            continue
        if kind == "none":
            average = midpoint_average(flairs)
            if earlier.get(user) == REASON_AVERAGE:
                value, reason, extra = birth_year_value((average, average)), REASON_AVERAGE, None
                outcome["average (as before)"] += 1
            else:
                value = NEEDS_DECISION
                reason = (
                    "No common year, not even by first appearance; decide by hand from columns D/E "
                    f"(average of the midpoints would be {average})"
                )
                extra = (flairs_text(flairs), ranges_summary(flairs))
                outcome["needs decision"] += 1
            new_rows.append([user, value, reason, *(extra or (None, None))])
            status_of[user] = (NEEDS_DECISION, span)
            continue
        reason = {"intersection": REASON_INTERSECTION, "crossing": REASON_CROSSING, "first": REASON_FIRST}[kind]
        new_rows.append([user, birth_year_value(birth_range), reason, None, None])
        outcome[kind] += 1
        status_of[user] = (
            {"intersection": "resolved", "crossing": "resolved (majority bucket)", "first": "resolved (first appearance)"}[kind],
            span,
        )

    decisions.delete_rows(2, decisions.max_row)
    for row in [*kept, *new_rows]:
        decisions.append(row)

    del workbook["Age report"]
    report = workbook.create_sheet("Age report")
    report.append(REPORT_HEADER)
    for user in sorted(status_of):
        status, span = status_of[user]
        for (subreddit, flair), years in sorted(users[user].items()):
            annotation = annotations[(subreddit, flair)]
            report.append([
                user,
                status,
                span,
                subreddit,
                flair,
                years_posted(years),
                *(annotation.get(column) for column in ANNOTATION_COLUMNS),
                ranges_text(years),
            ])
    workbook.save(audit_path)

    no_longer = sorted(set(earlier) - contradictory)
    print(f"{len(contradictory):,} users with contradictory age comments from 2012 on")
    for name, count in sorted(outcome.items()):
        print(f"  {name}: {count:,}")
    print(f"{len(no_longer):,} users with an earlier rule-based decision no longer need one")
    hand_idle = sorted(set(decided_by_hand) - contradictory)
    if hand_idle:
        print(f"Decided by hand but no longer contradictory (decision kept): {', '.join(hand_idle)}")
    for user, hand, rule in hand_differs:
        print(f"Decided by hand as {hand!r}; the rules would now give {rule!r}: {user}")
    for row in new_rows:
        if row[1] == NEEDS_DECISION:
            print(f"Needs decision: {row[0]} ({row[2]})")
    print(f"Decisions and Age report -> {audit_path} (backup in {BACKUP_DIR.name}/)")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT, help="3a age.jsonl")
    parser.add_argument("--flairs", type=Path, default=GROUPED_FLAIRS_PATH)
    parser.add_argument("--audit", type=Path, default=AUDIT_PATH)
    args = parser.parse_args()
    run(args.input, args.flairs, args.audit)


if __name__ == "__main__":
    main()
